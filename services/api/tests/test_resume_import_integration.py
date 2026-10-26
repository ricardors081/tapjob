"""Integration tests: real PostgreSQL, full HTTP lifecycle for WP1.3
(docs/adr/0011-resume-import.md). Worker processing is invoked directly
(`service.process_resume_import`) rather than through a real Celery broker -- this test suite
proves the DOMAIN behavior the task wraps, not Celery/Redis transport itself (which has no
business logic of its own here).
"""

import io
import threading
import uuid
import zipfile
from http import HTTPStatus
from typing import Any, cast

import docx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from identity_fakes import FakeIdentityVerifier, bearer_header, make_identity
from sqlalchemy import create_engine, select

from tapjob_api.career_profile.models import ProfileFact
from tapjob_api.core.config import ApiSettings
from tapjob_api.db.session import create_session_factory
from tapjob_api.identity.models import AccountAuditEvent
from tapjob_api.resume_import import service
from tapjob_api.resume_import.models import (
    CandidateReviewState,
    ImportFailureCode,
    ImportStatus,
    Resume,
    ResumeImport,
    ResumeImportCandidate,
    ResumeVersion,
)
from tapjob_api.resume_import.parsing import DeterministicResumeParser
from tapjob_api.resume_import.storage import ObjectStoragePort

pytestmark = pytest.mark.integration

BOOTSTRAP_PATH = "/api/v1/auth/bootstrap"
IMPORTS_PATH = "/api/v1/me/resume-imports"
MASTER_RESUME_PATH = "/api/v1/me/resumes/master"


def _resume_storage(client: TestClient) -> ObjectStoragePort:
    """`TestClient.app`'s static type is a bare ASGI callable, not `FastAPI` -- this is the one,
    type-checked place that narrows it to read back the SAME storage adapter instance the app's
    own lifespan constructed (`tapjob_api.main`), rather than repeating an unchecked attribute
    access at every call site."""
    app = client.app
    assert isinstance(app, FastAPI)
    return cast(ObjectStoragePort, app.state.resume_storage)


def _docx_resume_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Summary")
    document.add_paragraph(
        "Senior backend engineer focused on distributed systems and developer platforms."
    )
    document.add_paragraph("Experience")
    document.add_paragraph("Senior Backend Engineer at Globex Corp")
    document.add_paragraph("Jan 2022 - Present")
    document.add_paragraph("Led migration to event-driven architecture.")
    document.add_paragraph("Skills")
    document.add_paragraph("Python, PostgreSQL, Kubernetes")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _malformed_zip_bytes() -> bytes:
    return b"PK\x03\x04" + b"not a real zip central directory" * 3


def _bootstrap_and_upload(
    client: TestClient, verifier: FakeIdentityVerifier, *, token: str, subject: str
) -> tuple[uuid.UUID, uuid.UUID]:
    """Returns (user_id, import_id)."""
    verifier.issue_token(token, make_identity(subject=subject))
    bootstrap = client.post(BOOTSTRAP_PATH, headers=bearer_header(token))
    assert bootstrap.status_code == HTTPStatus.OK
    user_id = uuid.UUID(bootstrap.json()["user"]["id"])

    response = client.post(
        IMPORTS_PATH,
        headers=bearer_header(token),
        files={
            "file": (
                "resume.docx",
                _docx_resume_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == HTTPStatus.ACCEPTED
    import_id = uuid.UUID(response.json()["id"])
    return user_id, import_id


def test_full_lifecycle_upload_review_finalize_creates_master_resume(
    identity_client: TestClient, fake_identity_verifier: FakeIdentityVerifier
) -> None:
    settings = ApiSettings()
    # The API process (identity_client's app) constructed its OWN InMemoryObjectStorage in its
    # lifespan; this test reuses that SAME instance (via `app.state.resume_storage`) to simulate
    # the worker directly, proving the real end-to-end object-storage round trip without a real
    # Celery/Redis broker.
    engine = create_engine(str(settings.database_url))
    session_factory = create_session_factory(engine)

    user_id, import_id = _bootstrap_and_upload(
        identity_client, fake_identity_verifier, token="token-1", subject="uid-1"
    )

    with session_factory() as session:
        resume_import = session.get(ResumeImport, import_id)
        assert resume_import is not None
        assert resume_import.status == ImportStatus.QUEUED

        # Zero ProfileFacts exist before any processing/review/finalize (docs/adr/0011-...md "core
        # product invariant").
        assert session.scalar(select(ProfileFact).where(ProfileFact.user_id == user_id)) is None

    # Simulate the worker using the identity_client app's OWN storage adapter (reachable via its
    # `app.state`), proving the real end-to-end object-storage round trip.
    storage = _resume_storage(identity_client)
    with session_factory() as session:
        service.process_resume_import(
            session, import_id=import_id, storage=storage, parser=DeterministicResumeParser()
        )

    with session_factory() as session:
        resume_import = session.get(ResumeImport, import_id)
        assert resume_import is not None
        assert resume_import.status == ImportStatus.READY_FOR_REVIEW
        assert resume_import.parser_identity is not None
        candidates = list(
            session.scalars(
                select(ResumeImportCandidate).where(
                    ResumeImportCandidate.resume_import_id == import_id
                )
            )
        )
        assert len(candidates) > 0
        assert all(c.review_state == CandidateReviewState.PENDING for c in candidates)
        # Still zero ProfileFacts -- candidates are reviewable, not confirmed.
        assert session.scalar(select(ProfileFact).where(ProfileFact.user_id == user_id)) is None

    get_response = identity_client.get(
        f"{IMPORTS_PATH}/{import_id}", headers=bearer_header("token-1")
    )
    assert get_response.status_code == HTTPStatus.OK
    body = get_response.json()
    assert body["status"] == "ready_for_review"
    assert body["candidates"] is not None
    candidate_ids = [c["id"] for c in body["candidates"]]
    skill_candidate = next(c for c in body["candidates"] if c["candidateKind"] == "skill")

    decisions = [
        {"candidateId": cid, "decision": "accept"}
        if cid != skill_candidate["id"]
        else {
            "candidateId": cid,
            "decision": "edit_accept",
            "editedValue": {"kind": "skill", "name": "Golang"},
        }
        for cid in candidate_ids
    ]
    finalize_response = identity_client.post(
        f"{IMPORTS_PATH}/{import_id}/finalize",
        headers=bearer_header("token-1"),
        json={"expectedProfileVersion": 0, "decisions": decisions},
    )
    assert finalize_response.status_code == HTTPStatus.OK
    finalize_body = finalize_response.json()
    assert finalize_body["version"] == 1
    assert finalize_body["profileVersion"] == 1
    # Every candidate is accepted (none rejected) -- `factIds` excludes exactly the headline and
    # summary candidates, since those become CareerProfile top-level fields, never ProfileFacts
    # (docs/adr/0011-...md "career profile top-level fields").
    assert len(finalize_body["factIds"]) == len(candidate_ids) - 2

    with session_factory() as session:
        facts = list(session.scalars(select(ProfileFact).where(ProfileFact.user_id == user_id)))
        assert len(facts) > 0
        for fact in facts:
            assert fact.source_kind.value == "resume_import"
            assert fact.source_ref == str(import_id)
        skill_facts = [f for f in facts if f.fact_type.value == "skill"]
        assert any(f.value_json.get("name") == "Golang" for f in skill_facts)

        resume_import = session.get(ResumeImport, import_id)
        assert resume_import is not None
        assert resume_import.status == ImportStatus.FINALIZED
        assert resume_import.finalized_resume_id is not None

        resume = session.get(Resume, resume_import.finalized_resume_id)
        assert resume is not None
        assert resume.is_master is True

    master_response = identity_client.get(MASTER_RESUME_PATH, headers=bearer_header("token-1"))
    assert master_response.status_code == HTTPStatus.OK
    master_body = master_response.json()
    assert master_body["version"] == 1
    assert master_body["profileVersion"] == 1
    assert set(master_body["factIds"]) == set(finalize_body["factIds"])

    # Audit events never carry resume content.
    with session_factory() as session:
        events = list(
            session.scalars(
                select(AccountAuditEvent).where(
                    AccountAuditEvent.user_id == user_id,
                    AccountAuditEvent.event_type == "resume_import_finalized",
                )
            )
        )
        assert len(events) == 1
        metadata_blob = str(events[0].metadata)
        assert "Globex" not in metadata_blob
        assert "distributed systems" not in metadata_blob

    engine.dispose()


def test_finalize_is_idempotent_on_retry(
    identity_client: TestClient, fake_identity_verifier: FakeIdentityVerifier
) -> None:
    settings = ApiSettings()
    engine = create_engine(str(settings.database_url))
    session_factory = create_session_factory(engine)

    user_id, import_id = _bootstrap_and_upload(
        identity_client, fake_identity_verifier, token="token-2", subject="uid-2"
    )
    storage = _resume_storage(identity_client)
    with session_factory() as session:
        service.process_resume_import(
            session, import_id=import_id, storage=storage, parser=DeterministicResumeParser()
        )

    get_response = identity_client.get(
        f"{IMPORTS_PATH}/{import_id}", headers=bearer_header("token-2")
    )
    decisions = [
        {"candidateId": c["id"], "decision": "reject"} for c in get_response.json()["candidates"]
    ]
    first = identity_client.post(
        f"{IMPORTS_PATH}/{import_id}/finalize",
        headers=bearer_header("token-2"),
        json={"expectedProfileVersion": 0, "decisions": decisions},
    )
    assert first.status_code == HTTPStatus.OK

    second = identity_client.post(
        f"{IMPORTS_PATH}/{import_id}/finalize",
        headers=bearer_header("token-2"),
        json={"expectedProfileVersion": 0, "decisions": decisions},
    )
    assert second.status_code == HTTPStatus.OK
    assert second.json() == first.json()

    with session_factory() as session:
        resume_count = len(list(session.scalars(select(Resume).where(Resume.user_id == user_id))))
        version_count = len(
            list(
                session.scalars(select(ResumeVersion).join(Resume).where(Resume.user_id == user_id))
            )
        )
        assert resume_count == 1
        assert version_count == 1
    engine.dispose()


def test_worker_redelivery_of_an_already_processed_import_never_duplicates_candidates(
    identity_client: TestClient, fake_identity_verifier: FakeIdentityVerifier
) -> None:
    """Celery's own delivery-safety defaults (`task_acks_late`/`task_reject_on_worker_lost`) mean
    the SAME task message can be delivered more than once -- `process_resume_import` must treat a
    second delivery of an already-`PROCESSING`/`READY_FOR_REVIEW` import as a safe no-op, never a
    second round of candidate inserts (docs/adr/0011-resume-import.md "worker idempotency")."""
    settings = ApiSettings()
    engine = create_engine(str(settings.database_url))
    session_factory = create_session_factory(engine)

    _user_id, import_id = _bootstrap_and_upload(
        identity_client, fake_identity_verifier, token="token-retry", subject="uid-retry"
    )
    storage = _resume_storage(identity_client)

    with session_factory() as session:
        service.process_resume_import(
            session, import_id=import_id, storage=storage, parser=DeterministicResumeParser()
        )
    with session_factory() as session:
        # Simulates Celery redelivering the exact same task message after the first delivery
        # already completed -- `claim_import_for_processing`'s conditional UPDATE finds the row no
        # longer `QUEUED` and this call returns immediately without touching candidates at all.
        service.process_resume_import(
            session, import_id=import_id, storage=storage, parser=DeterministicResumeParser()
        )

    with session_factory() as session:
        resume_import = session.get(ResumeImport, import_id)
        assert resume_import is not None
        assert resume_import.status == ImportStatus.READY_FOR_REVIEW
        candidates = list(
            session.scalars(
                select(ResumeImportCandidate).where(
                    ResumeImportCandidate.resume_import_id == import_id
                )
            )
        )
        candidate_ids = [c.id for c in candidates]
        assert len(candidate_ids) == len(set(candidate_ids))  # no duplicates
        assert len(candidates) > 0
    engine.dispose()


def test_concurrent_finalizations_of_the_same_import_produce_exactly_one_logical_finalization(
    identity_client: TestClient, fake_identity_verifier: FakeIdentityVerifier
) -> None:
    """Several concurrent `finalize` calls for the SAME import, identical decisions/expected
    version, race on `finalize_resume_import`'s own `SELECT ... FOR UPDATE` -- exactly one of them
    performs the real write (`is_replay=False`), every other one observes the already-`FINALIZED`
    row and replays the SAME result (docs/adr/0011-...md "finalization idempotency",
    "concurrency"). Never two `Resume`/`ResumeVersion` rows, never two `profile_version` bumps."""
    settings = ApiSettings()
    engine = create_engine(str(settings.database_url))
    session_factory = create_session_factory(engine)

    user_id, import_id = _bootstrap_and_upload(
        identity_client, fake_identity_verifier, token="token-concurrent", subject="uid-concurrent"
    )
    storage = _resume_storage(identity_client)
    with session_factory() as session:
        service.process_resume_import(
            session, import_id=import_id, storage=storage, parser=DeterministicResumeParser()
        )

    get_response = identity_client.get(
        f"{IMPORTS_PATH}/{import_id}", headers=bearer_header("token-concurrent")
    )
    decisions = [
        {"candidateId": c["id"], "decision": "reject"} for c in get_response.json()["candidates"]
    ]
    payload = {"expectedProfileVersion": 0, "decisions": decisions}

    results: list[dict[str, Any]] = []
    lock = threading.Lock()

    def _call() -> None:
        response = identity_client.post(
            f"{IMPORTS_PATH}/{import_id}/finalize",
            headers=bearer_header("token-concurrent"),
            json=payload,
        )
        with lock:
            results.append({"status": response.status_code, "body": response.json()})

    threads = [threading.Thread(target=_call) for _ in range(6)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()

    assert all(result["status"] == HTTPStatus.OK for result in results)
    resume_version_ids = {result["body"]["resumeVersionId"] for result in results}
    assert len(resume_version_ids) == 1  # every caller observed the SAME finalization

    with session_factory() as session:
        resume_count = len(list(session.scalars(select(Resume).where(Resume.user_id == user_id))))
        version_count = len(
            list(
                session.scalars(select(ResumeVersion).join(Resume).where(Resume.user_id == user_id))
            )
        )
        assert resume_count == 1
        assert version_count == 1
    engine.dispose()


def test_stale_profile_version_at_finalize_is_a_clean_conflict_with_zero_partial_writes(
    identity_client: TestClient, fake_identity_verifier: FakeIdentityVerifier
) -> None:
    settings = ApiSettings()
    engine = create_engine(str(settings.database_url))
    session_factory = create_session_factory(engine)

    user_id, import_id = _bootstrap_and_upload(
        identity_client, fake_identity_verifier, token="token-3", subject="uid-3"
    )
    storage = _resume_storage(identity_client)
    with session_factory() as session:
        service.process_resume_import(
            session, import_id=import_id, storage=storage, parser=DeterministicResumeParser()
        )

    # The profile changes BETWEEN parsing and finalization.
    patch_response = identity_client.patch(
        "/api/v1/me/career-profile",
        headers=bearer_header("token-3"),
        json={"expectedProfileVersion": 0, "headline": "Unrelated manual edit"},
    )
    assert patch_response.status_code == HTTPStatus.OK
    assert patch_response.json()["profileVersion"] == 1

    get_response = identity_client.get(
        f"{IMPORTS_PATH}/{import_id}", headers=bearer_header("token-3")
    )
    decisions = [
        {"candidateId": c["id"], "decision": "reject"} for c in get_response.json()["candidates"]
    ]
    finalize_response = identity_client.post(
        f"{IMPORTS_PATH}/{import_id}/finalize",
        headers=bearer_header("token-3"),
        json={"expectedProfileVersion": 0, "decisions": decisions},  # stale on purpose
    )
    assert finalize_response.status_code == HTTPStatus.CONFLICT
    assert finalize_response.json()["code"] == "PROFILE_VERSION_CONFLICT"

    with session_factory() as session:
        resume_import = session.get(ResumeImport, import_id)
        assert resume_import is not None
        assert resume_import.status == ImportStatus.READY_FOR_REVIEW  # unchanged
        assert session.scalar(select(Resume).where(Resume.user_id == user_id)) is None
        candidates = list(
            session.scalars(
                select(ResumeImportCandidate).where(
                    ResumeImportCandidate.resume_import_id == import_id
                )
            )
        )
        assert all(c.review_state == CandidateReviewState.PENDING for c in candidates)
    engine.dispose()


def test_user_b_cannot_read_review_or_finalize_user_as_resume_import(
    identity_client: TestClient, fake_identity_verifier: FakeIdentityVerifier
) -> None:
    fake_identity_verifier.issue_token("token-a", make_identity(subject="uid-a"))
    bootstrap_a = identity_client.post(BOOTSTRAP_PATH, headers=bearer_header("token-a"))
    assert bootstrap_a.status_code == HTTPStatus.OK
    upload = identity_client.post(
        IMPORTS_PATH,
        headers=bearer_header("token-a"),
        files={
            "file": (
                "resume.docx",
                _docx_resume_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == HTTPStatus.ACCEPTED
    import_id = upload.json()["id"]

    fake_identity_verifier.issue_token("token-b", make_identity(subject="uid-b"))
    bootstrap_b = identity_client.post(BOOTSTRAP_PATH, headers=bearer_header("token-b"))
    assert bootstrap_b.status_code == HTTPStatus.OK

    read_as_b = identity_client.get(f"{IMPORTS_PATH}/{import_id}", headers=bearer_header("token-b"))
    assert read_as_b.status_code == HTTPStatus.NOT_FOUND
    assert read_as_b.json()["code"] == "RESUME_IMPORT_NOT_FOUND"

    finalize_as_b = identity_client.post(
        f"{IMPORTS_PATH}/{import_id}/finalize",
        headers=bearer_header("token-b"),
        json={"expectedProfileVersion": 0, "decisions": []},
    )
    assert finalize_as_b.status_code in (HTTPStatus.NOT_FOUND, HTTPStatus.UNPROCESSABLE_ENTITY)


def test_an_unsupported_upload_is_rejected_with_a_normalized_error_code(
    identity_client: TestClient, fake_identity_verifier: FakeIdentityVerifier
) -> None:
    fake_identity_verifier.issue_token("token-4", make_identity(subject="uid-4"))
    bootstrap = identity_client.post(BOOTSTRAP_PATH, headers=bearer_header("token-4"))
    assert bootstrap.status_code == HTTPStatus.OK

    response = identity_client.post(
        IMPORTS_PATH,
        headers=bearer_header("token-4"),
        files={"file": ("resume.txt", b"plain text, not a real document" * 5, "text/plain")},
    )
    assert response.status_code == HTTPStatus.UNPROCESSABLE_ENTITY
    assert response.json()["code"] == "UNSUPPORTED_DOCUMENT_TYPE"


def test_a_malformed_document_fails_processing_with_a_deterministic_failed_state(
    identity_client: TestClient, fake_identity_verifier: FakeIdentityVerifier
) -> None:
    settings = ApiSettings()
    engine = create_engine(str(settings.database_url))
    session_factory = create_session_factory(engine)

    fake_identity_verifier.issue_token("token-5", make_identity(subject="uid-5"))
    bootstrap = identity_client.post(BOOTSTRAP_PATH, headers=bearer_header("token-5"))
    assert bootstrap.status_code == HTTPStatus.OK

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("_rels/.rels", "<Relationships/>")
        archive.writestr("word/document.xml", "not valid docx XML content at all")
    upload = identity_client.post(
        IMPORTS_PATH,
        headers=bearer_header("token-5"),
        files={
            "file": (
                "resume.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == HTTPStatus.ACCEPTED
    import_id = uuid.UUID(upload.json()["id"])

    storage = _resume_storage(identity_client)
    with session_factory() as session:
        service.process_resume_import(
            session, import_id=import_id, storage=storage, parser=DeterministicResumeParser()
        )

    with session_factory() as session:
        resume_import = session.get(ResumeImport, import_id)
        assert resume_import is not None
        assert resume_import.status == ImportStatus.FAILED
        assert resume_import.failure_code is not None
    engine.dispose()


def test_a_document_with_meaningful_text_but_no_recognizable_candidates_fails_cleanly(
    identity_client: TestClient, fake_identity_verifier: FakeIdentityVerifier
) -> None:
    """Round-1 independent review "Finding 8": text extraction can succeed while the deterministic
    parser recognizes zero reviewable candidates (no section headers, no plausible headline, no
    links). Without an explicit domain outcome for this, the import would reach
    `READY_FOR_REVIEW` with zero candidates -- a state `FinalizeResumeImportCommand.decisions`'s own
    `min_length=1` makes impossible to ever finalize. Must instead be a normalized, terminal
    `FAILED` with `NO_REVIEWABLE_CANDIDATES` -- never `INTERNAL_ERROR`, never `READY_FOR_REVIEW`."""
    settings = ApiSettings()
    engine = create_engine(str(settings.database_url))
    session_factory = create_session_factory(engine)

    fake_identity_verifier.issue_token("token-6", make_identity(subject="uid-6"))
    bootstrap = identity_client.post(BOOTSTRAP_PATH, headers=bearer_header("token-6"))
    assert bootstrap.status_code == HTTPStatus.OK
    user_id = uuid.UUID(bootstrap.json()["user"]["id"])

    document = docx.Document()
    # A single long preamble line (over the headline length bound) with no recognizable section
    # header, no email marker, no year, no URL -- genuinely zero candidates, not merely a short
    # fixture that happens to dodge every heuristic by accident.
    document.add_paragraph(
        "This introductory paragraph intentionally contains no recognizable resume section "
        "headings, no plausible short headline, no email address, no year, and no link, so "
        "that the deterministic parser recognizes exactly zero reviewable candidates from it "
        "despite real, meaningful, successfully extracted text content being present throughout."
    )
    buffer = io.BytesIO()
    document.save(buffer)
    upload = identity_client.post(
        IMPORTS_PATH,
        headers=bearer_header("token-6"),
        files={
            "file": (
                "resume.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == HTTPStatus.ACCEPTED
    import_id = uuid.UUID(upload.json()["id"])

    storage = _resume_storage(identity_client)
    with session_factory() as session:
        service.process_resume_import(
            session, import_id=import_id, storage=storage, parser=DeterministicResumeParser()
        )

    with session_factory() as session:
        resume_import = session.get(ResumeImport, import_id)
        assert resume_import is not None
        assert resume_import.status == ImportStatus.FAILED
        assert resume_import.failure_code == ImportFailureCode.NO_REVIEWABLE_CANDIDATES
        assert (
            session.scalar(
                select(ResumeImportCandidate).where(
                    ResumeImportCandidate.resume_import_id == import_id
                )
            )
            is None
        )
        assert session.scalar(select(ProfileFact).where(ProfileFact.user_id == user_id)) is None
        assert session.scalar(select(Resume).where(Resume.user_id == user_id)) is None
    engine.dispose()
