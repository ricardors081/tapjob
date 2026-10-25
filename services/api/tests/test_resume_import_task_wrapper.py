"""Tests of the Celery TASK WRAPPER itself (`tasks.process_resume_import_task`), not just
`service.process_resume_import` in isolation -- round-2 independent review "Finding 1"/"Finding 3"
explicitly required this: the service layer raising the right exception is necessary but not
sufficient if the Celery wrapper around it computes the wrong retry countdown/budget.

Calls `process_resume_import_task.run(...)` directly (Celery's own documented unit-testing pattern
-- `bind=True` makes `.run` a method ALREADY bound to the real, singleton task/app-proxy object, so
there is no separate "self" to fake; instead, `self.retry` is monkeypatched directly on that real
task object, exactly like `test_resume_import_dispatch_recovery.py` already overrides
`deps.get_task_producer` via `app.dependency_overrides` for the same reason -- a recording double
in place of the real side-effecting call). This never touches a real broker/Redis.
"""

import io
import uuid
from datetime import UTC, datetime, timedelta
from http import HTTPStatus
from typing import Any, cast

import docx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from identity_fakes import FakeIdentityVerifier, bearer_header, make_identity
from sqlalchemy import create_engine, update

from tapjob_api.core.config import ApiSettings
from tapjob_api.db.session import create_session_factory
from tapjob_api.resume_import import service, tasks
from tapjob_api.resume_import.errors import ImportLeaseActiveError, TransientProcessingError
from tapjob_api.resume_import.limits import (
    LEASE_WAIT_MAX_RETRIES,
    LEASE_WAIT_MIN_COUNTDOWN_SECONDS,
    LEASE_WAIT_SAFETY_MARGIN_SECONDS,
    MAX_PROCESSING_ATTEMPTS,
    PROCESSING_LEASE_SECONDS,
    RESUME_PROCESS_HARD_TIME_LIMIT_SECONDS,
    RESUME_PROCESS_SOFT_TIME_LIMIT_SECONDS,
    TRANSIENT_RETRY_COUNTDOWN_SECONDS,
)
from tapjob_api.resume_import.models import ImportStatus, ResumeImport
from tapjob_api.resume_import.storage import ObjectStoragePort

pytestmark = pytest.mark.integration

BOOTSTRAP_PATH = "/api/v1/auth/bootstrap"
IMPORTS_PATH = "/api/v1/me/resume-imports"


class _RetrySentinelError(Exception):
    """Raised by the fake `retry` double in place of Celery's real `Retry`/
    `MaxRetriesExceededError` -- lets a test assert on exactly what `self.retry(...)` was called
    with, without needing a real Celery broker/eager-mode machinery."""

    def __init__(self, *, exc: Exception, countdown: float, max_retries: int | None) -> None:
        self.exc = exc
        self.countdown = countdown
        self.max_retries = max_retries
        super().__init__(f"retry(countdown={countdown}, max_retries={max_retries})")


def _fake_retry(*, exc: Exception, countdown: float, max_retries: int | None = None) -> Exception:
    return _RetrySentinelError(exc=exc, countdown=countdown, max_retries=max_retries)


def _resume_storage(client: TestClient) -> ObjectStoragePort:
    app = client.app
    assert isinstance(app, FastAPI)
    return cast(ObjectStoragePort, app.state.resume_storage)


# -- pure countdown computation -------------------------------------------------------------------


def test_lease_wait_retry_countdown_is_computed_from_the_leases_real_remaining_lifetime() -> None:
    now = datetime(2026, 1, 1, tzinfo=UTC)
    retry_at = now + timedelta(seconds=120)

    countdown = tasks.lease_wait_retry_countdown(retry_at, now=now)

    assert countdown == pytest.approx(120 + LEASE_WAIT_SAFETY_MARGIN_SECONDS)


def test_lease_wait_retry_countdown_has_a_floor_even_when_the_lease_already_expired() -> None:
    now = datetime(2026, 1, 1, tzinfo=UTC)
    retry_at = now - timedelta(seconds=999)  # already long expired

    countdown = tasks.lease_wait_retry_countdown(retry_at, now=now)

    assert countdown == LEASE_WAIT_MIN_COUNTDOWN_SECONDS


# -- soft/hard time limit invariants --------------------------------------------------------------


def test_soft_and_hard_time_limits_are_configured_and_strictly_bounded_by_the_lease() -> None:
    """Round-2 independent review "Finding 3": `soft < hard < lease` is a real, asserted invariant
    -- not merely documentation -- so a genuinely still-running worker is PROVABLY hard-killed
    before its own lease could ever expire out from under it."""
    assert (
        tasks.process_resume_import_task.soft_time_limit == RESUME_PROCESS_SOFT_TIME_LIMIT_SECONDS
    )
    assert tasks.process_resume_import_task.time_limit == RESUME_PROCESS_HARD_TIME_LIMIT_SECONDS
    assert RESUME_PROCESS_SOFT_TIME_LIMIT_SECONDS < RESUME_PROCESS_HARD_TIME_LIMIT_SECONDS
    assert RESUME_PROCESS_HARD_TIME_LIMIT_SECONDS < PROCESSING_LEASE_SECONDS


def test_a_soft_time_limit_during_processing_fails_the_row_terminally_without_a_retry_loop(
    identity_client: TestClient,
    fake_identity_verifier: FakeIdentityVerifier,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """`celery.exceptions.SoftTimeLimitExceeded` is a plain `Exception` subclass in this pinned
    Celery version (confirmed by direct inspection) -- already caught by
    `service.process_resume_import`'s own generic exception handler exactly like any other
    unexpected failure: terminal `FAILED`, correctly lease-fenced, then re-raised for observability
    (never automatically retried -- a bounded, resource-capped parsing job timing out is itself an
    anomalous condition, not a presumed-transient one)."""
    from celery.exceptions import SoftTimeLimitExceeded

    settings = ApiSettings()
    engine = create_engine(str(settings.database_url))
    session_factory = create_session_factory(engine)

    fake_identity_verifier.issue_token("timeout-1", make_identity(subject="uid-timeout-1"))
    bootstrap = identity_client.post(BOOTSTRAP_PATH, headers=bearer_header("timeout-1"))
    assert bootstrap.status_code == HTTPStatus.OK
    upload = identity_client.post(
        IMPORTS_PATH,
        headers=bearer_header("timeout-1"),
        files={
            "file": (
                "resume.pdf",
                b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj\n<< >>\nendobj\n%%EOF",
                "application/pdf",
            )
        },
    )
    assert upload.status_code == HTTPStatus.ACCEPTED
    import_id = uuid.UUID(upload.json()["id"])
    storage = _resume_storage(identity_client)

    def _raise_soft_time_limit(*_: Any, **__: Any) -> str:
        raise SoftTimeLimitExceeded()

    monkeypatch.setattr(service, "extract_pdf_text", _raise_soft_time_limit)

    with session_factory() as session, pytest.raises(SoftTimeLimitExceeded):
        service.process_resume_import(session, import_id=import_id, storage=storage)

    with session_factory() as session:
        resume_import = session.get(ResumeImport, import_id)
        assert resume_import is not None
        assert resume_import.status == ImportStatus.FAILED
        assert resume_import.processing_lease_token is None
        assert resume_import.processing_attempts == 1  # never retried again
    engine.dispose()


# -- the task wrapper itself ------------------------------------------------------------------


def test_task_wrapper_retries_a_transient_failure_with_the_tasks_own_default_budget(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(tasks.process_resume_import_task, "retry", _fake_retry)
    monkeypatch.setattr(tasks, "_session_factory", lambda: _NullSession())
    monkeypatch.setattr(tasks, "_storage", object())

    def _raise_transient(session: Any, *, import_id: Any, storage: Any, **kwargs: Any) -> None:
        raise TransientProcessingError("boom")

    monkeypatch.setattr(service, "process_resume_import", _raise_transient)

    with pytest.raises(_RetrySentinelError) as exc_info:
        tasks.process_resume_import_task.run(str(uuid.uuid4()))

    assert exc_info.value.countdown == TRANSIENT_RETRY_COUNTDOWN_SECONDS
    # No override: this call relies on the TASK's own configured `max_retries`
    # (`MAX_PROCESSING_ATTEMPTS`), the real, DB-tracked processing-attempt budget.
    assert exc_info.value.max_retries is None


def test_task_wrapper_retries_an_active_lease_with_a_countdown_past_expiry_and_a_separate_budget(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Round-2 independent review "Finding 1": the task wrapper itself -- not just the service
    layer -- must compute its countdown from the lease's own real remaining lifetime, and must use
    a SEPARATE, larger retry budget than genuine processing attempts."""
    monkeypatch.setattr(tasks.process_resume_import_task, "retry", _fake_retry)
    monkeypatch.setattr(tasks, "_session_factory", lambda: _NullSession())
    monkeypatch.setattr(tasks, "_storage", object())

    retry_at = datetime.now(UTC) + timedelta(seconds=250)

    def _raise_lease_active(session: Any, *, import_id: Any, storage: Any, **kwargs: Any) -> None:
        raise ImportLeaseActiveError(retry_at=retry_at)

    monkeypatch.setattr(service, "process_resume_import", _raise_lease_active)

    with pytest.raises(_RetrySentinelError) as exc_info:
        tasks.process_resume_import_task.run(str(uuid.uuid4()))

    # Past the lease's own real expiry (250s) plus the safety margin -- never a fixed, lease-
    # duration-agnostic countdown. A tiny tolerance absorbs the real wall-clock time that elapses
    # between constructing `retry_at` above and `lease_wait_retry_countdown`'s own `now()` call.
    assert exc_info.value.countdown >= 250 + LEASE_WAIT_SAFETY_MARGIN_SECONDS - 1
    assert exc_info.value.max_retries == LEASE_WAIT_MAX_RETRIES
    assert exc_info.value.max_retries != MAX_PROCESSING_ATTEMPTS


class _NullSession:
    """A no-op context manager standing in for a real DB session -- these tests monkeypatch
    `service.process_resume_import` itself, so no real session is ever used."""

    def __enter__(self) -> None:
        return None

    def __exit__(self, *_: object) -> None:
        return None


# -- the full worker-crash-and-recovery flow the review explicitly asked for ----------------------


def test_worker_crash_after_claim_schedules_a_retry_past_the_leases_expiry_and_recovers(
    identity_client: TestClient,
    fake_identity_verifier: FakeIdentityVerifier,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The exact scenario round-2 independent review "Finding 1" asked for a deterministic test
    of: worker A claims at t0 and crashes (never completing); an IMMEDIATE redelivery must still
    schedule its retry for AFTER the real lease expiry, never a fixed guess; once the lease
    genuinely expires, a later attempt reclaims and completes; the import never remains stuck
    `PROCESSING`."""
    settings = ApiSettings()
    engine = create_engine(str(settings.database_url))
    session_factory = create_session_factory(engine)

    fake_identity_verifier.issue_token("timeout-2", make_identity(subject="uid-timeout-2"))
    bootstrap = identity_client.post(BOOTSTRAP_PATH, headers=bearer_header("timeout-2"))
    assert bootstrap.status_code == HTTPStatus.OK
    document = docx.Document()
    document.add_paragraph("Summary")
    document.add_paragraph("Experienced backend engineer.")
    document.add_paragraph("Skills")
    document.add_paragraph("Python, PostgreSQL")
    buffer = io.BytesIO()
    document.save(buffer)
    upload = identity_client.post(
        IMPORTS_PATH,
        headers=bearer_header("timeout-2"),
        files={
            "file": (
                "resume.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == HTTPStatus.ACCEPTED
    import_id = uuid.UUID(upload.json()["id"])
    storage = _resume_storage(identity_client)

    # Worker A claims the row and "crashes" -- this session's own claim is committed, but nothing
    # further ever runs in it (simulating a process death right after the claim).
    with session_factory() as session:
        claim_result = service.claim_import_for_processing(session, import_id=import_id)
        session.commit()
        assert claim_result is not None
        real_lease_expiry = claim_result[0].processing_lease_expires_at
    assert real_lease_expiry is not None

    # Immediate redelivery: the task wrapper observes the still-active lease and must schedule its
    # own retry for AFTER the real expiry, not a fixed guess.
    monkeypatch.setattr(tasks.process_resume_import_task, "retry", _fake_retry)
    monkeypatch.setattr(tasks, "_session_factory", lambda: session_factory())
    monkeypatch.setattr(tasks, "_storage", storage)

    with pytest.raises(_RetrySentinelError) as exc_info:
        tasks.process_resume_import_task.run(str(import_id))

    now = datetime.now(UTC)
    remaining_real_lease_seconds = (real_lease_expiry - now).total_seconds()
    assert exc_info.value.countdown >= remaining_real_lease_seconds
    assert exc_info.value.max_retries == LEASE_WAIT_MAX_RETRIES

    with session_factory() as session:
        resume_import = session.get(ResumeImport, import_id)
        assert resume_import is not None
        assert resume_import.status == ImportStatus.PROCESSING  # untouched by the retry itself
        assert resume_import.processing_attempts == 1

    # Now simulate real elapsed time past the lease's own expiry (exactly what the computed
    # countdown above waited for) and let a later attempt genuinely reclaim and complete.
    with session_factory() as session:
        session.execute(
            update(ResumeImport)
            .where(ResumeImport.id == import_id)
            .values(processing_lease_expires_at=datetime.now(UTC) - timedelta(seconds=1))
        )
        session.commit()
    with session_factory() as session:
        service.process_resume_import(session, import_id=import_id, storage=storage)

    with session_factory() as session:
        resume_import = session.get(ResumeImport, import_id)
        assert resume_import is not None
        # Never remains stuck PROCESSING -- the review's own explicit success criterion.
        assert resume_import.status == ImportStatus.READY_FOR_REVIEW
    engine.dispose()
