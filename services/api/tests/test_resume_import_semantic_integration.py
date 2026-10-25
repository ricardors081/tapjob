import io
import uuid
from collections.abc import Iterator
from datetime import UTC, datetime
from typing import Any

import docx
import pytest
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

from tapjob_api.core.config import ApiSettings
from tapjob_api.db.session import create_session_factory
from tapjob_api.identity.models import User
from tapjob_api.ports._common import (
    AdapterHealth,
    AdapterHealthStatus,
    CallContext,
    PortError,
    PortErrorCode,
)
from tapjob_api.ports.ai_gateway import (
    AiGatewayCapability,
    AiRequest,
    AiResult,
    AiSafetyStatus,
    AiUsage,
)
from tapjob_api.ports.capabilities import CapabilityManifest, CapabilityReport, CapabilityStatus
from tapjob_api.resume_import.models import ImportStatus
from tapjob_api.resume_import.parsing import PARSER_IDENTITY
from tapjob_api.resume_import.semantic_parser import SEMANTIC_PARSER_IDENTITY, SemanticResumeParser
from tapjob_api.resume_import.semantic_schema import (
    SemanticExperience,
    SemanticHeadline,
    SemanticResumeExtraction,
    SemanticSkill,
)
from tapjob_api.resume_import.service import (
    create_resume_import,
    process_resume_import,
)
from tapjob_api.resume_import.storage import InMemoryObjectStorage

pytestmark = pytest.mark.integration


class FlakyAiGateway:
    def __init__(
        self,
        *,
        should_fail: bool = False,
        failure_code: PortErrorCode = PortErrorCode.UNAVAILABLE,
    ) -> None:
        self.provider = "flaky-openai"
        self.should_fail = should_fail
        self.failure_code = failure_code

    async def capabilities(self, ctx: CallContext) -> CapabilityManifest[AiGatewayCapability]:
        return CapabilityManifest(
            [
                CapabilityReport(
                    capability=AiGatewayCapability.STRUCTURED_OUTPUT,
                    status=CapabilityStatus.SUPPORTED,
                )
            ]
        )

    async def health(self, ctx: CallContext) -> AdapterHealth:
        return AdapterHealth(status=AdapterHealthStatus.HEALTHY, checked_at=datetime.now(UTC))

    async def generate[OutputT](
        self, request: AiRequest[OutputT], ctx: CallContext
    ) -> AiResult[OutputT]:
        if self.should_fail:
            raise PortError(
                self.failure_code,
                provider=self.provider,
                message=f"Simulated error with code {self.failure_code}",
            )
        extraction = SemanticResumeExtraction(
            candidate_name="Julia Teste",
            professional_headline=SemanticHeadline(
                headline="Data Scientist", source_excerpt="Data Scientist"
            ),
            experiences=[
                SemanticExperience(
                    employer="Alpha Analytics",
                    title="Data Scientist",
                    technologies=["Python", "SQL"],
                    source_excerpt="Alpha Analytics",
                )
            ],
            skills=[SemanticSkill(name="Python", source_excerpt="Python")],
        )
        output = request.validate_output(extraction.model_dump(mode="json"))
        return AiResult(
            output=output,
            usage=AiUsage(
                input_tokens=100, output_tokens=50, latency_ms=25.0, estimated_cost_micros=None
            ),
            safety_status=AiSafetyStatus.ALLOWED,
            prompt=request.prompt,
        )


def _make_valid_docx() -> bytes:
    doc = docx.Document()
    doc.add_paragraph("Julia Teste")
    doc.add_paragraph("Data Scientist")
    doc.add_paragraph("Experiência:")
    doc.add_paragraph("Alpha Analytics — Data Scientist (2020 - Atual)")
    doc.add_paragraph("Atuação em modelos preditivos com Python e SQL.")
    doc.add_paragraph("Competências:")
    doc.add_paragraph("Python, SQL")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def db_session(_identity_schema_ready: None) -> Iterator[Session]:
    settings = ApiSettings()
    engine = create_engine(str(settings.database_url))
    factory = create_session_factory(engine)
    session = factory()
    try:
        yield session
    finally:
        session.close()
        with engine.begin() as connection:
            connection.execute(
                text("TRUNCATE TABLE resume_import_candidates, resume_imports, users CASCADE")
            )
        engine.dispose()


def test_process_resume_import_success_with_semantic_parser(db_session: Session) -> None:
    user = User(id=uuid.uuid4(), locale="pt-BR")
    db_session.add(user)
    db_session.commit()

    storage = InMemoryObjectStorage()
    imp = create_resume_import(
        db_session,
        user=user,
        content=_make_valid_docx(),
        original_filename="curriculo.docx",
        declared_content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage=storage,
    )

    gateway = FlakyAiGateway(should_fail=False)
    semantic_parser = SemanticResumeParser(gateway)

    process_resume_import(
        db_session,
        import_id=imp.id,
        storage=storage,
        parser=semantic_parser,
    )

    db_session.refresh(imp)
    assert imp.status == ImportStatus.READY_FOR_REVIEW
    assert imp.parser_identity == SEMANTIC_PARSER_IDENTITY
    assert len(imp.candidates) > 0


def test_process_resume_import_refusal_does_not_trigger_fallback(db_session: Session) -> None:
    """Requirement 3: Model refusal (INVALID_REQUEST) must NOT trigger deterministic fallback."""
    user = User(id=uuid.uuid4(), locale="pt-BR")
    db_session.add(user)
    db_session.commit()

    storage = InMemoryObjectStorage()
    imp = create_resume_import(
        db_session,
        user=user,
        content=_make_valid_docx(),
        original_filename="curriculo.docx",
        declared_content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage=storage,
    )

    gateway = FlakyAiGateway(should_fail=True, failure_code=PortErrorCode.INVALID_REQUEST)
    parser = SemanticResumeParser(gateway)

    with pytest.raises(PortError) as exc_info:
        process_resume_import(
            db_session,
            import_id=imp.id,
            storage=storage,
            parser=parser,
        )
    assert exc_info.value.code == PortErrorCode.INVALID_REQUEST

    db_session.refresh(imp)
    assert imp.status == ImportStatus.FAILED
    assert len(imp.candidates) == 0


def test_process_resume_import_auth_failure_does_not_trigger_fallback(db_session: Session) -> None:
    """Requirement 3: Auth failure (AUTH_FAILED) must NOT trigger deterministic fallback."""
    user = User(id=uuid.uuid4(), locale="pt-BR")
    db_session.add(user)
    db_session.commit()

    storage = InMemoryObjectStorage()
    imp = create_resume_import(
        db_session,
        user=user,
        content=_make_valid_docx(),
        original_filename="curriculo.docx",
        declared_content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage=storage,
    )

    gateway = FlakyAiGateway(should_fail=True, failure_code=PortErrorCode.AUTH_FAILED)
    parser = SemanticResumeParser(gateway)

    with pytest.raises(PortError) as exc_info:
        process_resume_import(
            db_session,
            import_id=imp.id,
            storage=storage,
            parser=parser,
        )
    assert exc_info.value.code == PortErrorCode.AUTH_FAILED

    db_session.refresh(imp)
    assert imp.status == ImportStatus.FAILED
    assert len(imp.candidates) == 0


def test_process_resume_import_timeout_triggers_fallback(db_session: Session) -> None:
    """Requirement 3: Transient timeout (TIMEOUT) triggers deterministic fallback."""
    user = User(id=uuid.uuid4(), locale="pt-BR")
    db_session.add(user)
    db_session.commit()

    storage = InMemoryObjectStorage()
    imp = create_resume_import(
        db_session,
        user=user,
        content=_make_valid_docx(),
        original_filename="curriculo.docx",
        declared_content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage=storage,
    )

    gateway = FlakyAiGateway(should_fail=True, failure_code=PortErrorCode.TIMEOUT)
    parser = SemanticResumeParser(gateway)

    process_resume_import(
        db_session,
        import_id=imp.id,
        storage=storage,
        parser=parser,
    )

    db_session.refresh(imp)
    assert imp.status == ImportStatus.READY_FOR_REVIEW
    assert imp.parser_identity == PARSER_IDENTITY
    assert len(imp.candidates) > 0


def test_process_resume_import_unavailable_triggers_fallback(db_session: Session) -> None:
    """Requirement 3: Transient unavailability (UNAVAILABLE) triggers deterministic fallback."""
    user = User(id=uuid.uuid4(), locale="pt-BR")
    db_session.add(user)
    db_session.commit()

    storage = InMemoryObjectStorage()
    imp = create_resume_import(
        db_session,
        user=user,
        content=_make_valid_docx(),
        original_filename="curriculo.docx",
        declared_content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage=storage,
    )

    gateway = FlakyAiGateway(should_fail=True, failure_code=PortErrorCode.UNAVAILABLE)
    parser = SemanticResumeParser(gateway)

    process_resume_import(
        db_session,
        import_id=imp.id,
        storage=storage,
        parser=parser,
    )

    db_session.refresh(imp)
    assert imp.status == ImportStatus.READY_FOR_REVIEW
    assert imp.parser_identity == PARSER_IDENTITY
    assert len(imp.candidates) > 0


def test_process_resume_import_does_not_catch_type_error(db_session: Session) -> None:
    """Hardening 5: Unexpected exceptions (TypeError, AttributeError)
    must NEVER trigger fallback."""
    user = User(id=uuid.uuid4(), locale="pt-BR")
    db_session.add(user)
    db_session.commit()

    storage = InMemoryObjectStorage()
    imp = create_resume_import(
        db_session,
        user=user,
        content=_make_valid_docx(),
        original_filename="curriculo.docx",
        declared_content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage=storage,
    )

    class BuggyParser:
        provider = "buggy-parser"

        async def parse(self, doc: Any, ctx: Any) -> Any:
            raise TypeError("Simulated internal programming bug in parser")

    with pytest.raises(TypeError, match="Simulated internal programming bug"):
        process_resume_import(
            db_session,
            import_id=imp.id,
            storage=storage,
            parser=BuggyParser(),  # type: ignore[arg-type]
        )


def test_process_resume_import_preserves_ambiguities(db_session: Session) -> None:
    """Hardening 6: Ambiguities from extraction are preserved in ResumeImport review context."""
    user = User(id=uuid.uuid4(), locale="pt-BR")
    db_session.add(user)
    db_session.commit()

    storage = InMemoryObjectStorage()
    imp = create_resume_import(
        db_session,
        user=user,
        content=_make_valid_docx(),
        original_filename="curriculo.docx",
        declared_content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage=storage,
    )

    from tapjob_api.resume_import.semantic_schema import SemanticAmbiguity

    class AmbiguousAiGateway:
        provider = "ambiguous-gateway"

        async def capabilities(self, ctx: CallContext) -> CapabilityManifest[AiGatewayCapability]:
            return CapabilityManifest(
                [
                    CapabilityReport(
                        capability=AiGatewayCapability.STRUCTURED_OUTPUT,
                        status=CapabilityStatus.SUPPORTED,
                    )
                ]
            )

        async def health(self, ctx: CallContext) -> AdapterHealth:
            return AdapterHealth(status=AdapterHealthStatus.HEALTHY, checked_at=datetime.now(UTC))

        async def generate[OutputT](
            self, request: AiRequest[OutputT], ctx: CallContext
        ) -> AiResult[OutputT]:
            extraction = SemanticResumeExtraction(
                candidate_name="Julia Teste",
                professional_headline=SemanticHeadline(
                    headline="Data Scientist", source_excerpt="Data Scientist"
                ),
                experiences=[
                    SemanticExperience(
                        employer="Alpha Analytics",
                        title="Data Scientist",
                        current=True,
                        source_excerpt="Alpha Analytics",
                    )
                ],
                ambiguities=[
                    SemanticAmbiguity(
                        field="graduation_year",
                        question="Qual foi o ano exato de conclusão do curso?",
                        options=["2018", "2019", "2020"],
                        source_excerpt="Concluído",
                    )
                ],
            )
            output = request.validate_output(extraction.model_dump(mode="json"))
            return AiResult(
                output=output,
                usage=AiUsage(
                    input_tokens=100,
                    output_tokens=50,
                    latency_ms=25.0,
                    estimated_cost_micros=None,
                ),
                safety_status=AiSafetyStatus.ALLOWED,
                prompt=request.prompt,
            )

    semantic_parser = SemanticResumeParser(AmbiguousAiGateway())
    process_resume_import(
        db_session,
        import_id=imp.id,
        storage=storage,
        parser=semantic_parser,
    )

    db_session.refresh(imp)
    assert imp.status == ImportStatus.READY_FOR_REVIEW
    assert imp.ambiguities_count == 1
    assert len(imp.ambiguities) == 1
    assert imp.ambiguities[0]["field"] == "graduation_year"
    assert "ano exato" in imp.ambiguities[0]["question"]


def test_semantic_parser_fail_closed_grounding_drops_ungrounded_facts() -> None:
    """Blocker 2: When source_excerpt is not found in document text, candidate facts are DROPPED."""
    from tapjob_api.ports.resume_parser import ResumeDocument
    from tapjob_api.resume_import.semantic_schema import (
        SemanticEducation,
        SemanticExperience,
        SemanticHeadline,
        SemanticSkill,
        SemanticSummary,
    )

    class UngroundedAiGateway:
        provider = "ungrounded-gateway"

        async def capabilities(self, ctx: CallContext) -> CapabilityManifest[AiGatewayCapability]:
            return CapabilityManifest(
                [
                    CapabilityReport(
                        capability=AiGatewayCapability.STRUCTURED_OUTPUT,
                        status=CapabilityStatus.SUPPORTED,
                    )
                ]
            )

        async def health(self, ctx: CallContext) -> AdapterHealth:
            return AdapterHealth(status=AdapterHealthStatus.HEALTHY, checked_at=datetime.now(UTC))

        async def generate[OutputT](
            self, request: AiRequest[OutputT], ctx: CallContext
        ) -> AiResult[OutputT]:
            # Hallucinated excerpts not present anywhere in doc
            extraction = SemanticResumeExtraction(
                candidate_name="Julia Teste",
                professional_headline=SemanticHeadline(
                    headline="AI Super Architect",
                    source_excerpt="FABRICATED HEADLINE EXCERPT NOT IN DOCUMENT",
                ),
                professional_summary=SemanticSummary(
                    summary="Hallucinated summary about leading worldwide projects.",
                    source_excerpt="FABRICATED SUMMARY EXCERPT NOT IN DOCUMENT",
                ),
                experiences=[
                    SemanticExperience(
                        employer="Fake Corp",
                        title="Chief AI Officer",
                        current=True,
                        source_excerpt="FABRICATED EXPERIENCE EXCERPT NOT IN DOCUMENT",
                    )
                ],
                education=[
                    SemanticEducation(
                        institution="MIT Artificial",
                        degree="PhD",
                        source_excerpt="FABRICATED EDUCATION EXCERPT NOT IN DOCUMENT",
                    )
                ],
                skills=[
                    SemanticSkill(
                        name="Quantum Computing",
                        source_excerpt="FABRICATED SKILL EXCERPT NOT IN DOCUMENT",
                    )
                ],
            )
            output = request.validate_output(extraction.model_dump(mode="json"))
            return AiResult(
                output=output,
                usage=AiUsage(
                    input_tokens=100,
                    output_tokens=50,
                    latency_ms=25.0,
                    estimated_cost_micros=None,
                ),
                safety_status=AiSafetyStatus.ALLOWED,
                prompt=request.prompt,
            )

    parser = SemanticResumeParser(UngroundedAiGateway())
    doc = ResumeDocument(
        content=b"Apenas este texto simples no curriculo.",
        media_type="text/plain",
        filename="simples.txt",
    )
    ctx = CallContext.with_timeout(correlation_id="test-grounding", timeout_seconds=10.0)

    import asyncio

    parsed = asyncio.run(parser.parse(doc, ctx))

    # All ungrounded items must have been dropped (0 candidates reach Truth Store)
    msg = f"Expected 0 facts due to fail-closed grounding, got {len(parsed.facts)}"
    assert len(parsed.facts) == 0, msg


def test_eval_runner_returns_not_run_when_key_missing() -> None:
    """Blocker 4: Eval runner returns NOT_RUN_MISSING_SECRET when OPENAI_API_KEY is not set."""
    import asyncio
    import os
    import sys
    from pathlib import Path
    from unittest.mock import patch

    repo_root = Path(__file__).resolve().parents[3]
    if str(repo_root) not in sys.path:
        sys.path.insert(0, str(repo_root))

    from tools.evals.run_resume_semantic_evals import main as eval_main

    with patch.dict(os.environ, {"OPENAI_API_KEY": ""}, clear=False):
        exit_code = asyncio.run(eval_main())
        assert exit_code == 0
